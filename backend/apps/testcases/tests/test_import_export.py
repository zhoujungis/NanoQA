"""导入 / 导出（§61）：解析、字段映射、预览不落库、确认入库、非法行跳过、导出内容。"""
import csv
import io
import json

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.accounts.models import User
from apps.projects.models import Module, Tag
from apps.testcases.models import TestCase, TestCaseVersion
from apps.testcases.services.import_export import guess_mapping, parse_steps
from apps.testcases.tests.base import NanoTestCase

EXPORT_HEADERS = [
    "用例编号", "标题", "模块路径", "优先级", "类型", "状态",
    "前置条件", "步骤", "预期结果", "标签", "负责人",
]


def build_xlsx(headers, rows) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_csv(headers, rows, encoding="utf-8-sig") -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    return buffer.getvalue().encode(encoding)


class ImportExportFixtureMixin:
    """导入 / 导出接口地址与文件上传 helper。"""

    @property
    def export_url(self) -> str:
        return f"/api/v1/projects/{self.project.id}/export/testcases/"

    @property
    def import_url(self) -> str:
        return f"/api/v1/projects/{self.project.id}/import/testcases/"

    def upload(self, content: bytes, name="cases.xlsx", **extra):
        payload = {"file": SimpleUploadedFile(name, content)}
        payload.update(extra)
        return self.client.post(self.import_url, payload, format="multipart")

    def preview(self, content: bytes, name="cases.xlsx", mapping=None):
        extra = {"dry_run": "true"}
        if mapping is not None:
            extra["mapping"] = json.dumps(mapping, ensure_ascii=False)
        return self.upload(content, name, **extra)

    def confirm(self, content: bytes, name="cases.xlsx", mapping=None):
        extra = {"dry_run": "false"}
        if mapping is not None:
            extra["mapping"] = json.dumps(mapping, ensure_ascii=False)
        return self.upload(content, name, **extra)


class StepsAndMappingTests(NanoTestCase):
    """纯函数：步骤解析与表头自动猜测。"""

    def test_parse_steps_numbered_with_arrow(self):
        steps, errors = parse_steps("1. 打开登录页 → 登录页展示\n2、输入账号密码 -> 校验通过")
        self.assertEqual(errors, [])
        self.assertEqual(steps, [
            {"order": 1, "action": "打开登录页", "expected": "登录页展示"},
            {"order": 2, "action": "输入账号密码", "expected": "校验通过"},
        ])

    def test_parse_steps_fallback_expected_column(self):
        steps, errors = parse_steps("打开页面\n点击按钮", fallback_expected="操作成功")
        self.assertEqual(errors, [])
        self.assertEqual([step["expected"] for step in steps], ["操作成功", "操作成功"])

    def test_parse_steps_missing_expected_is_error(self):
        steps, errors = parse_steps("1. 打开页面")
        self.assertEqual(steps, [])
        self.assertEqual(len(errors), 1)
        self.assertIn("预期结果", errors[0])

    def test_parse_steps_empty_text(self):
        steps, errors = parse_steps("")
        self.assertEqual(steps, [])
        self.assertEqual(errors, ["步骤不能为空"])

    def test_guess_mapping_chinese_and_english_headers(self):
        headers = ["用例编号", "Case Title", "Module Path", "优先级", "Steps", "预期结果", "Owner", "备注列"]
        mapping = guess_mapping(headers)
        self.assertEqual(mapping["case_key"], "用例编号")
        self.assertEqual(mapping["title"], "Case Title")
        self.assertEqual(mapping["module"], "Module Path")
        self.assertEqual(mapping["priority"], "优先级")
        self.assertEqual(mapping["steps"], "Steps")
        self.assertEqual(mapping["expected_result"], "预期结果")
        self.assertEqual(mapping["owner"], "Owner")
        self.assertNotIn("description", mapping)


class ExportTests(ImportExportFixtureMixin, NanoTestCase):
    def _create_two_cases(self):
        first = self.create_case(
            title="登录成功", priority="P0", status="active", precondition="用户已注册", expected_result="进入首页",
            module_id=self.module.id, tag_ids=[self.tag_smoke.id], owner_id=self.owner.id,
            steps=[{"action": "打开登录页", "expected": "登录页展示"}, {"action": "输入密码", "expected": "校验通过"}],
        )
        second = self.create_case(title="下单失败", priority="P2", case_type="regression", status="draft")
        return first, second

    def test_export_csv_utf8_bom_and_content(self):
        first, _second = self._create_two_cases()
        response = self.client.get(self.export_url, {"format": "csv"})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn(f'filename="{self.project.key}_testcases_', response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"), "CSV 必须带 UTF-8-BOM")

        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0], EXPORT_HEADERS)
        self.assertEqual(len(rows), 3)
        # 默认按更新时间倒序，行序不固定，按编号定位
        first_row = next(row for row in rows[1:] if row[0] == first["case_key"])
        self.assertEqual(first_row[0], first["case_key"])
        self.assertEqual(first_row[1], "登录成功")
        self.assertEqual(first_row[2], "Auth/Login")
        self.assertEqual(first_row[3], "P0")
        self.assertEqual(first_row[4], "功能")
        self.assertEqual(first_row[5], "激活")
        self.assertEqual(first_row[6], "用户已注册")
        self.assertEqual(first_row[7], "1. 打开登录页 → 登录页展示\n2. 输入密码 → 校验通过")
        self.assertEqual(first_row[8], "进入首页")
        self.assertEqual(first_row[9], "smoke")
        self.assertEqual(first_row[10], "owner")

    def test_export_xlsx_bold_header_and_width(self):
        self._create_two_cases()
        response = self.client.get(self.export_url, {"format": "xlsx"})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn(
            "spreadsheetml", response["Content-Type"],
        )
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual([cell.value for cell in sheet[1]], EXPORT_HEADERS)
        self.assertTrue(all(cell.font.bold for cell in sheet[1]), "表头必须加粗")
        self.assertEqual(sheet.column_dimensions["A"].width, 18)
        self.assertEqual(sheet.max_row, 3)

    def test_export_filters_match_list_scope(self):
        self._create_two_cases()
        response = self.client.get(self.export_url, {"format": "csv", "priority": "P0"})
        self.assertEqual(response.status_code, 200, response.content)
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], "登录成功")

        response = self.client.get(self.export_url, {"format": "csv", "module_id": self.module.id})
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2], "Auth/Login")

    def test_export_invalid_format_rejected(self):
        response = self.client.get(self.export_url, {"format": "pdf"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["code"], 40006)

    def test_export_requires_read_permission(self):
        outsider = User.objects.create_user(username="outsider", email="out@example.com", password="pass12345")
        self.client.force_authenticate(outsider)
        response = self.client.get(self.export_url, {"format": "csv"})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["code"], 40001)

        self.client.force_authenticate(None)
        response = self.client.get(self.export_url, {"format": "csv"})
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json()["code"], 40002)

    def test_export_round_trip_rows_are_importable(self):
        """导出文件原样回导：编号命中 → 更新动作。"""
        first, _second = self._create_two_cases()
        response = self.client.get(self.export_url, {"format": "csv"})
        body = self.preview(response.content, name="round-trip.csv")
        self.assertEqual(body.status_code, 200, body.content)
        data = body.json()["data"]
        self.assertEqual(data["total"], 2)
        self.assertEqual([row["action"] for row in data["rows"]], ["update", "update"])


class ImportPreviewTests(ImportExportFixtureMixin, NanoTestCase):
    HEADERS = ["标题", "模块路径", "优先级", "类型", "状态", "前置条件", "步骤", "预期结果", "标签", "负责人"]

    def valid_rows(self):
        return [
            ["登录成功", "Auth/Login", "P0", "功能", "激活", "用户已注册",
             "1. 打开登录页 → 登录页展示\n2. 输入密码 → 校验通过", "进入首页", "smoke,新标签", "owner"],
            ["下单失败", "订单/支付", "P2", "回归", "草稿", "",
             "1. 提交订单 → 库存不足提示", "", "", ""],
        ]

    def test_preview_does_not_write_database(self):
        response = self.preview(build_xlsx(self.HEADERS, self.valid_rows()))
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["total"], 2)
        self.assertEqual(data["valid"], 2)
        self.assertEqual(data["invalid"], 0)
        self.assertEqual([row["action"] for row in data["rows"]], ["create", "create"])
        self.assertEqual(data["rows"][0]["data"]["steps"][0], {"order": 1, "action": "打开登录页", "expected": "登录页展示"})
        self.assertEqual(data["rows"][0]["data"]["module"], "Auth/Login")
        # 不落库
        self.assertEqual(TestCase.objects.count(), 0)
        self.assertEqual(TestCaseVersion.objects.count(), 0)
        self.assertEqual(Module.objects.filter(project=self.project).count(), 3)
        self.assertEqual(Tag.objects.filter(project=self.project).count(), 2)

    def test_preview_reports_invalid_rows_with_line_numbers(self):
        rows = [
            ["", "Auth", "P9", "神秘", "外星", "", "", "", "", ""],
            ["合法用例", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""],
            ["缺预期的步骤", "", "P1", "功能", "草稿", "", "打开页面", "", "", ""],
        ]
        response = self.preview(build_csv(self.HEADERS, rows), name="bad.csv")
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["total"], 3)
        self.assertEqual(data["valid"], 1)
        self.assertEqual(data["invalid"], 2)

        first_errors = data["rows"][0]["errors"]
        self.assertTrue(any("第 2 行：标题不能为空" == item for item in first_errors))
        self.assertTrue(any("优先级取值非法" in item for item in first_errors))
        self.assertTrue(any("类型取值非法" in item for item in first_errors))
        self.assertTrue(any("状态取值非法" in item for item in first_errors))
        self.assertTrue(any("步骤不能为空" in item for item in first_errors))
        self.assertIn("第 4 行", data["rows"][2]["errors"][0])
        self.assertIn("预期结果", data["rows"][2]["errors"][0])
        self.assertEqual(TestCase.objects.count(), 0)

    def test_preview_marks_duplicate_and_empty_rows_as_skip(self):
        rows = [
            ["重复用例", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""],
            ["重复用例", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""],
            ["", "", "", "", "", "", "", "", "", ""],
        ]
        response = self.preview(build_csv(self.HEADERS, rows), name="dup.csv")
        data = response.json()["data"]
        self.assertEqual([row["action"] for row in data["rows"]], ["create", "skip", "skip"])
        self.assertIn("重复", data["rows"][1]["errors"][0])
        self.assertIn("空行", data["rows"][2]["errors"][0])

    def test_preview_with_field_mapping(self):
        headers = ["Case Title", "My Module", "Level", "Steps"]
        rows = [["自定义标题", "模块A", "HIGH", "1. 打开页面 → 正常"]]
        mapping = {"title": "Case Title", "module": "My Module", "steps": "Steps"}
        response = self.preview(build_csv(headers, rows), name="mapped.csv", mapping=mapping)
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["mapping"]["title"], "Case Title")
        self.assertEqual(data["valid"], 1)
        self.assertEqual(data["rows"][0]["data"]["title"], "自定义标题")
        self.assertEqual(data["rows"][0]["data"]["module"], "模块A")
        # 未映射列不参与校验，Level 列被忽略
        self.assertEqual(data["rows"][0]["errors"], [])

    def test_parse_failure_returns_business_code(self):
        response = self.upload(b"this is not a workbook", name="broken.xlsx")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["code"], 40006)

        response = self.upload(build_csv(self.HEADERS, []), name="cases.txt")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["code"], 40006)

    def test_import_requires_create_permission(self):
        outsider = User.objects.create_user(username="outsider", email="out@example.com", password="pass12345")
        self.client.force_authenticate(outsider)
        response = self.preview(build_csv(self.HEADERS, self.valid_rows()), name="x.csv")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["code"], 40001)


class ImportConfirmTests(ImportExportFixtureMixin, NanoTestCase):
    HEADERS = ImportPreviewTests.HEADERS

    def test_confirm_creates_cases_with_keys_modules_tags_and_versions(self):
        rows = [
            ["登录成功", "订单/支付", "P0", "功能", "激活", "用户已注册",
             "1. 打开登录页 → 登录页展示\n2. 输入密码 → 校验通过", "进入首页", "smoke,新标签", "owner"],
            ["下单失败", "", "P2", "回归", "草稿", "", "1. 提交订单 → 库存不足提示", "", "", ""],
        ]
        response = self.confirm(build_xlsx(self.HEADERS, rows))
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["created"], 2)
        self.assertEqual(data["updated"], 0)
        self.assertEqual(data["skipped"], 0)
        self.assertEqual(data["failed"], [])

        cases = list(TestCase.objects.filter(project=self.project).order_by("id"))
        self.assertEqual([case.case_key for case in cases], ["SHOP-TC-000001", "SHOP-TC-000002"])
        first = cases[0]
        self.assertEqual(first.title, "登录成功")
        self.assertEqual(first.priority, "P0")
        self.assertEqual(first.case_type, "functional")
        self.assertEqual(first.status, "active")
        self.assertEqual(first.precondition, "用户已注册")
        self.assertEqual(first.current_version, 1)
        self.assertEqual(first.owner_id, self.owner.id)

        # 版本 1 快照必须同时创建（坑 4）
        version = TestCaseVersion.objects.get(test_case=first, version=1)
        self.assertEqual(version.change_summary, "初始版本")
        self.assertEqual(version.expected_result, "进入首页")
        self.assertEqual(version.steps[1], {"order": 2, "action": "输入密码", "expected": "校验通过"})

        # 模块按路径自动创建
        pay = Module.objects.get(project=self.project, name="支付")
        self.assertEqual(pay.parent.name, "订单")
        self.assertEqual(cases[0].module_id, pay.id)
        self.assertIsNone(cases[1].module_id)

        # 标签自动创建 + 已有标签复用
        self.assertTrue(Tag.objects.filter(project=self.project, name="新标签").exists())
        self.assertEqual(
            sorted(tag.name for tag in first.tags.all()), ["smoke", "新标签"],
        )

    def test_confirm_skips_invalid_rows_and_reports_them(self):
        rows = [
            ["", "", "", "", "", "", "", "", "", ""],
            ["合法用例", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""],
            ["非法优先级", "", "P9", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""],
        ]
        response = self.confirm(build_csv(self.HEADERS, rows), name="mix.csv")
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["created"], 1)
        self.assertEqual(data["skipped"], 1)
        self.assertEqual(len(data["failed"]), 1)
        self.assertEqual(data["failed"][0]["line"], 4)
        self.assertIn("第 4 行", data["failed"][0]["errors"][0])
        self.assertEqual(TestCase.objects.count(), 1)

    def test_confirm_keys_increment_and_import_twice_updates(self):
        rows = [["导入用例", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""]]
        first = self.confirm(build_csv(self.HEADERS, rows), name="one.csv")
        self.assertEqual(first.json()["data"]["created"], 1)

        rows2 = [
            ["导入用例", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""],
            ["第二个用例", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""],
        ]
        second = self.confirm(build_csv(self.HEADERS, rows2), name="two.csv")
        data = second.json()["data"]
        self.assertEqual(data["created"], 1)
        self.assertEqual(data["updated"], 1)

        keys = list(
            TestCase.objects.filter(project=self.project).order_by("id").values_list("case_key", flat=True)
        )
        self.assertEqual(keys, ["SHOP-TC-000001", "SHOP-TC-000002"])
        self.assertEqual(TestCase.objects.count(), 2)

    def test_confirm_updates_by_case_key_and_bumps_version(self):
        case = self.create_case(title="旧标题", priority="P1", steps=[{"action": "旧步骤", "expected": "旧预期"}])
        rows = [[case["case_key"], "新标题", "", "P0", "", "", "", "1. 新步骤 → 新预期", "新预期", "", ""]]
        response = self.confirm(build_csv(EXPORT_HEADERS, rows), name="update.csv")
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["updated"], 1)
        self.assertEqual(data["created"], 0)

        updated = TestCase.objects.get(pk=case["id"])
        self.assertEqual(updated.title, "新标题")
        self.assertEqual(updated.priority, "P0")
        # 空单元格保留原值
        self.assertEqual(updated.case_type, "functional")
        self.assertEqual(updated.current_version, 2)
        version = TestCaseVersion.objects.get(test_case=updated, version=2)
        self.assertEqual(version.change_summary, "导入更新")
        self.assertEqual(version.steps, [{"order": 1, "action": "新步骤", "expected": "新预期"}])

    def test_confirm_without_title_mapping_rejected(self):
        rows = [["只写步骤", "1. 打开页面 → 正常"]]
        response = self.confirm(
            build_csv(["备注", "步骤"], rows), name="no-title.csv", mapping={"steps": "步骤"}
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["code"], 40006)

    def test_csv_gb18030_encoding_supported(self):
        rows = [["中文标题", "", "P1", "功能", "草稿", "", "1. 打开页面 → 正常", "", "", ""]]
        response = self.confirm(
            build_csv(self.HEADERS, rows, encoding="gb18030"), name="gbk.csv"
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["data"]["created"], 1)
        self.assertEqual(TestCase.objects.get().title, "中文标题")
